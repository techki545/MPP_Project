from __future__ import annotations

import csv
from collections import defaultdict
from collections.abc import Iterator, Sequence
from pathlib import Path
import re

from openpyxl import load_workbook

from .models import DocumentRecord
from .normalization import (
    clean_text,
    normalize_doi,
    normalize_source_id,
    normalize_title,
    stable_document_id,
)


_CSV_COLUMNS = (
    "ID",
    "Author",
    "Publication Year",
    "Title",
    "Publication Title",
    "DOI",
    "Url",
    "Abstract Note",
    "Language",
)
_DOWNLOAD_FLAG_HEADERS = frozenset({"是否下载", "下载"})
_SOURCE_ID_HEADERS = frozenset({"id", "source id", "source_id", "item id", "item_id"})
_YEAR_PATTERN = re.compile(r"(?:19|20)\d{2}")


def _select_csv_encoding(path: Path) -> str:
    """Probe in bounded chunks so decoding fallback does not require loading the CSV."""
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            while handle.read(1024 * 1024):
                pass
    except UnicodeDecodeError:
        return "gb18030"
    return "utf-8-sig"


def _header_key(value: object) -> str:
    return clean_text(value).casefold()


def _header_indexes(headers: Sequence[object]) -> dict[str, int]:
    indexes: dict[str, int] = {}
    expected = {_header_key(column): column for column in _CSV_COLUMNS}
    for index, header in enumerate(headers):
        canonical = expected.get(_header_key(header))
        if canonical is not None and canonical not in indexes:
            indexes[canonical] = index
    return indexes


def _cell(row: Sequence[object], indexes: dict[str, int], column: str) -> str:
    index = indexes.get(column)
    if index is None or index >= len(row):
        return ""
    return clean_text(row[index])


def _authors(value: str) -> tuple[str, ...]:
    authors: list[str] = []
    for raw_author in re.split(r"[;\uFF1B]", value):
        author = clean_text(raw_author)
        if author:
            authors.append(author)
    return tuple(authors)


def _year(value: str) -> int | None:
    match = _YEAR_PATTERN.search(value)
    return int(match.group(0)) if match else None


def _document_from_row(
    row: Sequence[object],
    indexes: dict[str, int],
    source_row: int,
    source_id: str,
    normalized_source_id: str,
) -> DocumentRecord:
    authors = _authors(_cell(row, indexes, "Author"))
    year = _year(_cell(row, indexes, "Publication Year"))
    title = _cell(row, indexes, "Title")
    abstract = _cell(row, indexes, "Abstract Note")
    doi = _cell(row, indexes, "DOI")
    normalized_doi = normalize_doi(doi)
    return DocumentRecord(
        document_id=stable_document_id(
            normalized_doi,
            title or "source-" + source_id,
            authors[0] if authors else "",
            year,
        ),
        source_row=source_row,
        title=title,
        normalized_title=normalize_title(title),
        authors=authors,
        year=year,
        journal=_cell(row, indexes, "Publication Title"),
        doi=doi,
        normalized_doi=normalized_doi,
        abstract=abstract,
        language=_cell(row, indexes, "Language"),
        fulltext_status="low_information" if not title and not abstract else "missing",
        source_id=source_id,
        normalized_source_id=normalized_source_id,
        url=_cell(row, indexes, "Url"),
    )


def load_csv_document_items(path: Path) -> Iterator[tuple[str, DocumentRecord]]:
    """Stream normalized source IDs with their Zotero-style metadata records."""
    path = Path(path)
    encoding = _select_csv_encoding(path)
    # Some legacy exports are overwhelmingly gb18030 with isolated bad bytes.
    # The fallback stays gb18030 and replaces only those malformed byte sequences.
    errors = "replace" if encoding == "gb18030" else "strict"
    with path.open("r", encoding=encoding, errors=errors, newline="") as handle:
        reader = csv.reader(handle)
        headers = next(reader, None)
        if headers is None:
            return
        indexes = _header_indexes(headers)
        for source_row, row in enumerate(reader, start=2):
            if not any(clean_text(value) for value in row):
                continue
            source_id = _cell(row, indexes, "ID")
            normalized_source_id = normalize_source_id(source_id)
            yield normalized_source_id, _document_from_row(
                row, indexes, source_row, source_id, normalized_source_id
            )


def load_csv_documents(path: Path) -> Iterator[DocumentRecord]:
    """Stream Zotero-style MPP metadata rows as canonical document records."""
    for _, document in load_csv_document_items(path):
        yield document


def load_download_flags(paths: Sequence[Path]) -> dict[str, tuple[str, ...]]:
    """Read download flags from collaborators' workbooks in read-only mode."""
    flags_by_source_id: defaultdict[str, list[str]] = defaultdict(list)
    for path in paths:
        workbook = load_workbook(Path(path), read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                headers = next(rows, None)
                if headers is None:
                    continue
                source_index = next(
                    (
                        index
                        for index, header in enumerate(headers)
                        if _header_key(header) in _SOURCE_ID_HEADERS
                    ),
                    None,
                )
                flag_index = next(
                    (
                        index
                        for index, header in enumerate(headers)
                        if clean_text(header) in _DOWNLOAD_FLAG_HEADERS
                    ),
                    None,
                )
                if source_index is None or flag_index is None:
                    continue
                for row in rows:
                    if source_index >= len(row) or flag_index >= len(row):
                        continue
                    source_id = normalize_source_id(row[source_index])
                    flag = clean_text(row[flag_index])
                    if source_id and flag:
                        flags_by_source_id[source_id].append(flag)
        finally:
            workbook.close()
    return {key: tuple(values) for key, values in flags_by_source_id.items()}
