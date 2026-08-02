from pathlib import Path

from openpyxl import Workbook
import pytest

from knowledge_base.metadata_loader import (
    CsvDecodeWarning,
    load_csv_document_items,
    load_csv_documents,
    load_download_flags,
)
from knowledge_base.normalization import (
    clean_text,
    normalize_doi,
    normalize_source_id,
    normalize_title,
    stable_document_id,
    tokenize_for_fts,
)


def test_normalization_cleans_text_and_builds_stable_identifiers() -> None:
    assert clean_text("\ufeff  A\n B  ") == "A B"
    assert normalize_doi(" DOI: https://doi.org/10.1000/ABC. ") == "10.1000/abc"
    assert normalize_title("ＭＰＰ: Trial (I)") == "mpptriali"
    assert normalize_source_id(" 0012.0 ") == "12"
    assert stable_document_id("10.1/abc", "ignored", "ignored", 2025).startswith("doi-")
    assert stable_document_id("", "title", "Zhang", 2025) == stable_document_id(
        "", "title", "Zhang", 2025
    )
    assert "mpp" in tokenize_for_fts("MPP child pneumonia")


def test_load_csv_ignores_blank_and_duplicate_headers_and_builds_stable_id(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "metadata.csv"
    csv_path.write_text(
        "ID,Item Type,Author,,Publication Year,Title,Publication Title,DOI,Url,Abstract Note,Pages,Language,Author\n"
        "0012,journalArticle,Zhang,,2025, MPP Steroid Trial ,Journal,https://doi.org/10.1/ABC,https://x,Results,1-8,en,Ignored\n",
        encoding="utf-8-sig",
    )

    documents = list(load_csv_documents(csv_path))

    assert len(documents) == 1
    document = documents[0]
    assert document.source_row == 2
    assert document.source_id == "0012"
    assert document.normalized_source_id == "12"
    assert document.normalized_doi == "10.1/abc"
    assert document.document_id.startswith("doi-")
    assert document.title == "MPP Steroid Trial"
    assert document.authors == ("Zhang",)
    assert document.year == 2025
    assert document.journal == "Journal"
    assert document.abstract == "Results"
    assert document.language == "en"
    assert document.url == "https://x"


def test_load_csv_falls_back_to_gb18030_only_after_utf8_decode_failure(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "metadata-gb.csv"
    content = (
        "ID,Author,Publication Year,Title,Publication Title,DOI,Url,Abstract Note,Language\n"
        "7,Li,2024,儿童支原体肺炎,医学期刊,,,摘要内容,zh\n"
    )
    csv_path.write_bytes(content.encode("gb18030"))

    documents = list(load_csv_documents(csv_path))

    assert [(document.title, document.language) for document in documents] == [
        ("儿童支原体肺炎", "zh")
    ]


def test_load_csv_keeps_rows_with_an_isolated_invalid_byte_in_gb18030_export(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "metadata-corrupt-gb.csv"
    content = (
        "ID,Author,Publication Year,Title,Publication Title,DOI,Url,Abstract Note,Language\n"
        "8,Li,2024,儿童支原体肺炎,医学期刊,,,摘要内容,zh\n"
    ).encode("gb18030")
    csv_path.write_bytes(content.replace("摘要内容".encode("gb18030"), b"\x80"))

    diagnostics = []
    document = list(load_csv_documents(csv_path, diagnostics=diagnostics))[0]

    assert document.title == "儿童支原体肺炎"
    assert document.abstract == "�"
    assert diagnostics[0].affected_rows == (2,)
    assert diagnostics[0].replacement_count == 1
    assert diagnostics[0].omitted_row_count == 0


def test_load_csv_bounds_gb18030_replacement_diagnostic_rows(tmp_path: Path) -> None:
    csv_path = tmp_path / "metadata-many-corrupt-gb.csv"
    content = (
        "ID,Author,Publication Year,Title,Publication Title,DOI,Url,Abstract Note,Language\n"
        "8,Li,2024,儿童支原体肺炎,医学期刊,,,摘要内容,zh\n"
        "9,Li,2024,儿童支原体肺炎,医学期刊,,,摘要内容,zh\n"
        "10,Li,2024,儿童支原体肺炎,医学期刊,,,摘要内容,zh\n"
    ).encode("gb18030")
    csv_path.write_bytes(content.replace("摘要内容".encode("gb18030"), b"\x80"))

    diagnostics = []
    documents = list(
        load_csv_documents(csv_path, diagnostics=diagnostics, max_diagnostic_rows=2)
    )

    assert len(documents) == 3
    assert diagnostics[0].affected_rows == (2, 3)
    assert diagnostics[0].replacement_count == 3
    assert diagnostics[0].omitted_row_count == 1


def test_load_csv_emits_one_bounded_warning_without_diagnostic_container(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "metadata-warning-gb.csv"
    content = (
        "ID,Title,Abstract Note\n"
        "8,儿童支原体肺炎,摘要内容\n"
        "9,儿童支原体肺炎,摘要内容\n"
    ).encode("gb18030")
    csv_path.write_bytes(content.replace("摘要内容".encode("gb18030"), b"\x80"))

    with pytest.warns(CsvDecodeWarning) as caught:
        documents = list(load_csv_documents(csv_path))

    assert len(documents) == 2
    assert len(caught) == 1


def test_low_information_csv_rows_are_retained_but_marked_not_embeddable(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "low-information.csv"
    csv_path.write_text(
        "ID,Author,Publication Year,Title,Publication Title,DOI,Url,Abstract Note,Language\n"
        "99,Zhang,2025,,Journal,,,,en\n",
        encoding="utf-8",
    )

    document = next(iter(load_csv_documents(csv_path)))

    assert document.title == ""
    assert document.abstract == ""
    assert document.fulltext_status == "low_information"
    assert document.has_fulltext is False
    assert document.document_id.startswith("meta-")


def test_low_information_rows_without_metadata_keep_distinct_source_based_ids(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "empty-records.csv"
    csv_path.write_text(
        "ID,Author,Publication Year,Title,Publication Title,DOI,Url,Abstract Note,Language\n"
        "1,,,,,,,,\n"
        "2,,,,,,,,\n",
        encoding="utf-8",
    )

    documents = list(load_csv_documents(csv_path))

    assert len({document.document_id for document in documents}) == 2


def test_load_csv_document_items_returns_normalized_source_ids(tmp_path: Path) -> None:
    csv_path = tmp_path / "source-ids.csv"
    csv_path.write_text(
        "ID,Title,Abstract Note\n"
        " 0012.0 ,Trial,Abstract\n",
        encoding="utf-8",
    )

    source_id, document = next(iter(load_csv_document_items(csv_path)))

    assert source_id == "12"
    assert document.title == "Trial"


def _write_download_workbook(path: Path, flag_header: str, source_id: object, flag: object) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ID", "Title", flag_header])
    sheet.append([source_id, "Trial", flag])
    workbook.save(path)
    workbook.close()


def test_load_download_flags_streams_both_real_flag_header_variants(tmp_path: Path) -> None:
    yang = tmp_path / "yang.xlsx"
    han = tmp_path / "han.xlsx"
    _write_download_workbook(yang, "是否下载", " 0012 ", "是")
    _write_download_workbook(han, "下载", 12.0, "已下载")

    flags = load_download_flags((yang, han))

    assert flags == {"12": ("是", "已下载")}
